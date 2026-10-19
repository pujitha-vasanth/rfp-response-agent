"""Reads an RFP file (PDF, Word or text) and returns its text."""

from pathlib import Path

from docx import Document
from pypdf import PdfReader


def read_rfp(path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages)

    if suffix == ".docx":
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        # RFPs often put requirements in tables, so read those too.
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)

    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8")

    raise ValueError(f"Unsupported file type: {suffix}. Use PDF, DOCX or TXT.")
